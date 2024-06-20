import tkinter as tk
from tkinter import Listbox, END
from tkinter import ttk
from tkinter import filedialog
from tkinter import messagebox
from functools import partial
import pandas as pd
import numpy as np
from openpyxl import load_workbook
import calendar
import threading
import time
from functools import partial
import os
from categorize_date import categorize_date

processing = False
def get_family(row):
    key = (row['Division'], row['Article'])
    return Familly_Factors.get(key, '')

def process_File(filtered_mb51_df, nmclmc_df, ordre_df, multiplication_factors, category):
    merged_category_df = pd.merge(nmclmc_df, ordre_df, on="Nomenclature", how="left")
    
    for index, row in merged_category_df.iterrows():
        ordre = int(row['Ordre']) if not pd.isnull(row['Ordre']) else None
        composant = row['Component description']
        composant_unité = row['Unité de quantité']
        uq_de_saisie = row['Unité de qté de base']
        date_cat = category
        mb51qte_series = filtered_mb51_df.groupby('Ordre')['Quantité'].sum()
        mb51qte = mb51qte_series.get(ordre, default=0)

        if not filtered_mb51_df.empty:
            exact_quantity = abs(mb51_df[(mb51_df['Ordre'] == ordre) & 
                                        (mb51_df['Désignation article'] == composant) & (mb51_df['Date comptable'] == date_cat)]['Quantité'].sum())
            montant_di = mb51_df[(mb51_df['Ordre'] == ordre) & 
                                    (mb51_df['Désignation article'] == composant)]['Montant DI'].sum()
        else:
            
            mb51qte = 0
            exact_quantity = 0
            montant_di = 0
        
        
        if (row["Composant"], uq_de_saisie, composant_unité) in multiplication_factors:
            multiplication_factor = multiplication_factors[((row["Composant"]), uq_de_saisie, composant_unité)]
            exact_quantity *= multiplication_factor

        merged_category_df.at[index, 'mb51qte'] = mb51qte
        merged_category_df.at[index, 'exact_quantity'] = exact_quantity
        merged_category_df.at[index, 'Montant DI'] = montant_di  

    merged_category_df['Quantity Standard'] = abs((merged_category_df['Quantity'] * merged_category_df['mb51qte']) / merged_category_df['Quantité de base']).round(1)
    valid_rows = merged_category_df['Quantity Standard'] != 0
    merged_category_df.loc[valid_rows, 'Ecart'] = (((merged_category_df.loc[valid_rows, 'exact_quantity'] - merged_category_df.loc[valid_rows, 'Quantity Standard']) / merged_category_df.loc[valid_rows, 'Quantity Standard']) * 100).round(2)
    merged_category_df.loc[~valid_rows, 'Ecart'] = np.nan  
    merged_category_df['PROD Version'] = category

    return merged_category_df

def process_data_thread(root,input_file_entry, user_choice_entry,user_choice_combobox2, process_button, progress_bar):
    global processing

    if processing:
        messagebox.showinfo("Processing", "Data is still being processed. Please wait.")
        return

    input_file = input_file_entry.get()
    user_choice = user_choice_entry.get()
    user_choice2 = user_choice_combobox2.get()
    global df, nmclmc_df, ordre_df, mb51_df, df5, mat_type_df, df4, Famille_df,Familly_Factors,multiplication_factors,grouped_mb51
    process_button.config(state=tk.DISABLED)
    process_button.config(text="Processing...")


    processing = True

    try:
        messagebox.showinfo("Information", "Processing data. Please wait...")

        df = pd.read_excel(input_file, sheet_name='NMCL')
        nmclmc_df = pd.read_excel(input_file, sheet_name='NMCL')
        ordre_df = pd.read_excel(input_file, sheet_name="Ordre")
        mb51_df = pd.read_excel(input_file, sheet_name="Mb51")
        df5 = pd.read_excel(input_file, sheet_name='divdict')
        mat_type_df = pd.read_excel(input_file, sheet_name='MatType')
        df4 = pd.read_excel(input_file, sheet_name="Multiply")
        Famille_df = pd.read_excel(input_file, sheet_name='Family')

        
        grouped_mb51 = mb51_df.groupby(mb51_df["Date comptable"].apply(lambda x: categorize_date(x, user_choice)))
        mb51_df["Date comptable"] = grouped_mb51["Date comptable"].transform(lambda x: categorize_date(x.iloc[0], user_choice))
         
        multiplication_factors = df4.set_index(['Composant', 'Unité de Article', 'Unité de Composant']).to_dict()['multiplication factor']
        Familly_Factors = Famille_df.set_index(['Division', 'Article']).to_dict()['Famille']

        merged_df = pd.DataFrame()

        total_categories = len(grouped_mb51)
        processed_categories = 0

        for category, filtered_mb51_df in grouped_mb51:
            filtered_mb51_df = filtered_mb51_df[(filtered_mb51_df['Code mouvement'].isin([101, 102, 531, 532]))]
            processed_df = process_File(filtered_mb51_df, nmclmc_df, ordre_df, multiplication_factors, category)
            merged_df = pd.concat([merged_df, processed_df], ignore_index=True)
           
            processed_categories += 1
            progress_value = (processed_categories / total_categories) * 100
            progress_bar['value'] = progress_value
            root.update_idletasks()  # update the GUI
        
        
        merged_df['Montant DI'] = abs(merged_df['Montant DI'])
        valid_rows = merged_df['exact_quantity'] != 0
        merged_df.loc[valid_rows, 'Standard Montant DI'] = (((merged_df.loc[valid_rows, 'Montant DI'] / merged_df.loc[valid_rows, 'exact_quantity']) * merged_df.loc[valid_rows, 'Quantity Standard'])).round(2)
        merged_df['Ecart DI'] = (((merged_df['Montant DI'] - merged_df['Standard Montant DI']) / merged_df['Standard Montant DI']) * 100).round(2)
        
        # Merging data
        merged_df = pd.merge(merged_df, mat_type_df, left_on='Composant', right_on='Composant', how='left')
        merged_df= pd.merge(merged_df, df5, left_on='Division_x', right_on='Division', how='left')
        merged_df['Famille'] = merged_df.apply(get_family, axis=1)
        
        merged_df = merged_df[['Site','Division_x','Famille', 'Ordre','Article', 'Désignation article', 'MatTyp', 'Quantité de base', 'Unité de qté de base', 'Composant', 'Component description','Quantity', 'Unité de quantité', 'Nomenclature', 'PROD Version', 'mb51qte', 'Quantity Standard', 'exact_quantity', 'Ecart','Standard Montant DI','Montant DI','Ecart DI']]
        merged_df = merged_df.rename(columns={'Division_x': 'Division', 'mb51qte': 'Prod', "exact_quantity": "Cons ACT"})
        merged_df = merged_df.sort_values(by=["Division", "Article", "Nomenclature", "Ordre", "Prod","PROD Version"])
        merged_df = merged_df[merged_df['Prod'] != 0]

        cols_to_convert = ['Cons ACT', 'Quantity Standard', 'Montant DI', 'Standard Montant DI', 'Ecart DI', 'Ecart']
        merged_df[cols_to_convert] = merged_df[cols_to_convert].apply(pd.to_numeric, errors='coerce').fillna(0)

        

        
        if (user_choice == 'month'):
            merged_df1= merged_df[['Site','Division','Famille', 'Article', 'Désignation article', 'MatTyp', 'Quantité de base', 'Unité de qté de base', 'Composant', 'Component description','Quantity', 'Unité de quantité', 'Nomenclature', 'PROD Version', 'Prod', 'Quantity Standard', 'Cons ACT', 'Ecart']]
            merged_df1 = merged_df.groupby(['Site', 'Division', 'Famille', 'Article', 'Désignation article', 'MatTyp', 'Quantité de base', 'Unité de qté de base', 'Composant', 'Component description', 'Quantity', 'Unité de quantité', 'Nomenclature', 'PROD Version']).agg({'Prod': 'sum', 'Quantity Standard': 'sum', 'Cons ACT': 'sum', 'Ecart' : 'mean'}).reset_index()
            merged_df2= merged_df[['Site','Division','Famille', 'Article', 'Désignation article', 'MatTyp', 'Quantité de base', 'Unité de qté de base', 'Composant', 'Component description','Quantity', 'Unité de quantité', 'Nomenclature', 'PROD Version', 'Prod', 'Standard Montant DI', 'Montant DI', 'Ecart DI']]
            merged_df2 = merged_df.groupby(['Site', 'Division', 'Famille', 'Article', 'Désignation article', 'MatTyp', 'Quantité de base', 'Unité de qté de base', 'Composant', 'Component description', 'Quantity', 'Unité de quantité', 'Nomenclature', 'PROD Version']).agg({'Prod': 'sum','Standard Montant DI': 'sum', 'Montant DI': 'sum','Ecart DI' : 'mean'}).reset_index()


            
        else :
            merged_df1= merged_df[['Site','Division','Famille','Ordre', 'Article', 'Désignation article', 'MatTyp', 'Quantité de base', 'Unité de qté de base', 'Composant', 'Component description','Quantity', 'Unité de quantité', 'Nomenclature', 'PROD Version', 'Prod', 'Quantity Standard', 'Cons ACT', 'Ecart']]
            merged_df1 = merged_df.groupby(['Site', 'Division', 'Famille','Ordre', 'Article', 'Désignation article', 'MatTyp', 'Quantité de base', 'Unité de qté de base', 'Composant', 'Component description', 'Quantity', 'Unité de quantité', 'Nomenclature', 'PROD Version']).agg({'Prod': 'sum', 'Quantity Standard': 'sum', 'Cons ACT': 'sum','Ecart' : 'mean'}).reset_index()
            merged_df2= merged_df[['Site','Division','Famille','Ordre', 'Article', 'Désignation article', 'MatTyp', 'Quantité de base', 'Unité de qté de base', 'Composant', 'Component description','Quantity', 'Unité de quantité', 'Nomenclature', 'PROD Version', 'Prod', 'Standard Montant DI', 'Montant DI', 'Ecart DI']]
            merged_df2 = merged_df.groupby(['Site', 'Division', 'Famille','Ordre', 'Article', 'Désignation article', 'MatTyp', 'Quantité de base', 'Unité de qté de base', 'Composant', 'Component description', 'Quantity', 'Unité de quantité', 'Nomenclature', 'PROD Version']).agg({'Prod': 'sum','Standard Montant DI': 'sum', 'Montant DI': 'sum','Ecart DI' : 'mean'}).reset_index()

           
        ecart_by_mattype = merged_df1.groupby('MatTyp')['Ecart'].mean().reset_index()
        mattypdf = merged_df1.groupby('MatTyp')[['Cons ACT', 'Quantity Standard']].sum().reset_index()
        ecart_by_Div = merged_df1.groupby('Division')['Ecart'].mean().reset_index()
        Divdf = merged_df1.groupby('Division')[['Cons ACT', 'Quantity Standard']].sum().reset_index()
        ecart_comp =merged_df1.groupby('Composant')['Ecart'].mean().reset_index()
        qte_comp = merged_df1.groupby('Composant')[['Cons ACT', 'Quantity Standard']].sum().reset_index()
        ecart_Fam= merged_df1.groupby('Famille')['Ecart'].mean().reset_index()
        qtefam = merged_df1.groupby('Famille')[['Cons ACT', 'Quantity Standard']].sum().reset_index()
        ecart_dec= merged_df1.groupby('PROD Version')['Ecart'].mean().reset_index()
        qtedec = merged_df1.groupby('PROD Version')[['Cons ACT', 'Quantity Standard']].sum().reset_index()

        ecart_by_mattype1 = merged_df2.groupby('MatTyp')['Ecart DI'].mean().reset_index()
        mattypdf1 = merged_df2.groupby('MatTyp')[['Montant DI', 'Standard Montant DI']].sum().reset_index()
        ecart_by_Div1 = merged_df2.groupby('Division')['Ecart DI'].mean().reset_index()
        Divdf1 = merged_df2.groupby('Division')[['Montant DI', 'Standard Montant DI']].sum().reset_index()
        ecart_comp1 =merged_df2.groupby('Composant')['Ecart DI'].mean().reset_index()
        qte_comp1 = merged_df2.groupby('Composant')[['Montant DI', 'Standard Montant DI']].sum().reset_index()
        ecart_Fam1 = merged_df2.groupby('Famille')['Ecart DI'].mean().reset_index()
        qtefam1 = merged_df2.groupby('Famille')[['Montant DI', 'Standard Montant DI']].sum().reset_index()
        ecart_dec1 = merged_df2.groupby('PROD Version')['Ecart DI'].mean().reset_index()
        qtedec1 = merged_df2.groupby('PROD Version')[['Montant DI', 'Standard Montant DI']].sum().reset_index()
        
        merged_df = merged_df[merged_df['Prod'] != 0]
        negative_quantity_df = merged_df1[merged_df1['Quantity'] < 0]
        negative_quantity_df.loc[:, ['Cons ACT', 'Quantity Standard','Ecart','Standard Montant DI','Montant DI']] = '-'
        merged_df1.update(negative_quantity_df)
        merged_df1 = merged_df1.sort_values(by=['Division', 'Article','PROD Version'], ascending= True)
        merged_df2 = merged_df2.sort_values(by=['Division', 'Article','PROD Version'], ascending= True)
        
        
        if user_choice2 == 'STD VS Réel':
            result_file_final = "AnalysNMCL_STD_VS_Réel.xlsx"
            with pd.ExcelWriter(result_file_final, engine='xlsxwriter') as writer:
                merged_df1.to_excel(writer, sheet_name='Detailed Quantity Analysis', index=False)
                ecart_comp.to_excel(writer, sheet_name='Composant Analysis', index= False)
                qte_comp.to_excel(writer, sheet_name='Composant Analysis', startrow=0, startcol=13 , index=False)
                ecart_by_mattype.to_excel(writer, sheet_name='MatType Analysis', index=False)
                mattypdf.to_excel(writer, sheet_name='MatType Analysis', startrow=0, startcol=13 , index=False)
                ecart_by_Div.to_excel(writer, sheet_name='Division Analysis', index=False)
                Divdf.to_excel(writer, sheet_name='Division Analysis', startrow=0, startcol=13 , index=False)
                ecart_Fam.to_excel(writer, sheet_name='Family Analysis', index= False)
                qtefam.to_excel(writer, sheet_name='Family Analysis', startrow=0, startcol=13 , index=False)
                ecart_dec.to_excel(writer, sheet_name='Decade Analysis', index= False)
                qtedec.to_excel(writer, sheet_name='Decade Analysis', startrow=0, startcol=13 , index=False)
            
                workbook = writer.book
                
                worksheet = writer.sheets['MatType Analysis']
                
                chart_line_ecart = workbook.add_chart({'type': 'line'})
                chart_line_ecart.add_series({
                'categories': ['MatType Analysis', 1, 0, len(ecart_by_mattype), 0],  
                'values':     ['MatType Analysis', 1, 1, len(ecart_by_mattype), 1],  
                'name':       'Ecart by Mattype',
                })

                chart_line_ecart.width = 600  
                chart_line_ecart.height = 360
                worksheet.insert_chart('D1', chart_line_ecart)  
                chart_column = workbook.add_chart({'type': 'column'})
                chart_column.add_series({
                    'categories': ['MatType Analysis', 1, 0, len(mattypdf), 0],  
                    'values':     ['MatType Analysis', 1, 14, len(mattypdf), 14],  
                    'name':       'Cons ACT',
                })
                chart_column.add_series({
                    'values':     ['MatType Analysis', 1, 15, len(mattypdf), 15],  
                    'name':       'Quantity Standard',
                })
                chart_column.set_title({'name': 'Comparison of Cons ACT and Quantity Standard'})
                chart_column.set_x_axis({'name': 'MatTyp'})
                chart_column.set_y_axis({'name': 'Values'})
                
                chart_column.width = 600  
                chart_column.height = 360
                worksheet.insert_chart('R1', chart_column)


                worksheet = writer.sheets['Division Analysis']
                
                chart_line_ecart2 = workbook.add_chart({'type': 'line'})
                chart_line_ecart2.add_series({
                'categories': ['Division Analysis', 1, 0, len(ecart_by_Div), 0],  
                'values':     ['Division Analysis', 1, 1, len(ecart_by_Div), 1],  
                'name':       'Ecart by Division',
                })

                chart_line_ecart2.width = 600  
                chart_line_ecart2.height = 360
                worksheet.insert_chart('D1', chart_line_ecart2)  
                chart_column2 = workbook.add_chart({'type': 'column'})
                chart_column2.add_series({
                    'categories': ['Division Analysis', 1, 0, len(Divdf), 0],  
                    'values':     ['Division Analysis', 1, 14, len(Divdf), 14],  
                    'name':       'Cons ACT',
                })
                chart_column2.add_series({
                    'values':     ['Division Analysis', 1, 15, len(Divdf), 15],  
                    'name':       'Quantity Standard',
                })
                chart_column2.set_title({'name': 'Comparison of Cons ACT and Quantity Standard'})
                chart_column2.set_x_axis({'name': 'Division'})
                chart_column2.set_y_axis({'name': 'Values'})
                
                chart_column2.width = 600  
                chart_column2.height = 360
                worksheet.insert_chart('R1', chart_column2)

                worksheet = writer.sheets['Composant Analysis']
                
                chart_line_ecart3 = workbook.add_chart({'type': 'line'})
                chart_line_ecart3.add_series({
                'categories': ['Composant Analysis', 1, 0, len(ecart_comp), 0],  
                'values':     ['Composant Analysis', 1, 1, len(ecart_comp), 1],  
                'name':       'Ecart by Composant',
                })

                chart_line_ecart3.width = 600  
                chart_line_ecart3.height = 360
                worksheet.insert_chart('D1', chart_line_ecart3)  
                chart_column3 = workbook.add_chart({'type': 'column'})
                chart_column3.add_series({
                    'categories': ['Composant Analysis', 1, 0, len(qte_comp), 0],  
                    'values':     ['Composant Analysis', 1, 14, len(qte_comp), 14],  
                    'name':       'Cons ACT',
                })
                chart_column3.add_series({
                    'values':     ['Composant Analysis', 1, 15, len(qte_comp), 15],  
                    'name':       'Quantity Standard',
                })
                chart_column3.set_title({'name': 'Comparison of Cons ACT and Quantity Standard'})
                chart_column3.set_x_axis({'name': 'Composant'})
                chart_column3.set_y_axis({'name': 'Values'})
                
                chart_column3.width = 600  
                chart_column3.height = 360
                worksheet.insert_chart('R1', chart_column3)

                worksheet = writer.sheets['Family Analysis']
                
                chart_line_ecart4 = workbook.add_chart({'type': 'line'})
                chart_line_ecart4.add_series({
                'categories': ['Family Analysis', 1, 0, len(ecart_Fam), 0],  
                'values':     ['Family Analysis', 1, 1, len(ecart_Fam), 1],  
                'name':       'Ecart by Family',
                })

                chart_line_ecart4.width = 600  
                chart_line_ecart4.height = 360
                worksheet.insert_chart('D1', chart_line_ecart4)  
                chart_column4 = workbook.add_chart({'type': 'column'})
                chart_column4.add_series({
                    'categories': ['Family Analysis', 1, 0, len(qtefam), 0],  
                    'values':     ['Family Analysis', 1, 14, len(qtefam), 14],  
                    'name':       'Cons ACT',
                })
                chart_column4.add_series({
                    'values':     ['Family Analysis', 1, 15, len(qtefam), 15],  
                    'name':       'Quantity Standard',
                })
                chart_column4.set_title({'name': 'Comparison of Cons ACT and Quantity Standard'})
                chart_column4.set_x_axis({'name': 'Famille'})
                chart_column4.set_y_axis({'name': 'Values'})
                
                chart_column4.width = 600  
                chart_column4.height = 360
                worksheet.insert_chart('R1', chart_column4)

                worksheet = writer.sheets['Decade Analysis']
                
                chart_line_ecart5 = workbook.add_chart({'type': 'line'})
                chart_line_ecart5.add_series({
                'categories': ['Decade Analysis', 1, 0, len(ecart_dec), 0],  
                'values':     ['Decade Analysis', 1, 1, len(ecart_dec), 1],  
                'name':       'Ecart by Decade',
                })

                chart_line_ecart5.width = 600  
                chart_line_ecart5.height = 360
                worksheet.insert_chart('D1', chart_line_ecart5)  
                chart_column5 = workbook.add_chart({'type': 'column'})
                chart_column5.add_series({
                    'categories': ['Decade Analysis', 1, 0, len(qtedec), 0],  
                    'values':     ['Decade Analysis', 1, 14, len(qtedec), 14],  
                    'name':       'Cons ACT',
                })
                chart_column5.add_series({
                    'values':     ['Decade Analysis', 1, 15, len(qtedec), 15],  
                    'name':       'Quantity Standard',
                })
                chart_column5.set_title({'name': 'Comparison of Cons ACT and Quantity Standard'})
                chart_column5.set_x_axis({'name': 'Decade'})
                chart_column5.set_y_axis({'name': 'Values'})
                
                chart_column5.width = 600  
                chart_column5.height = 360
                worksheet.insert_chart('R1', chart_column5)
                
          

        elif user_choice2 == 'Valorisation':
            result_file_final = "AnalysNMCL_Valorisation.xlsx"
            with pd.ExcelWriter(result_file_final, engine='xlsxwriter') as writer:
                merged_df2.to_excel(writer, sheet_name='Detailed Valorisation Analysis', index=False)
                ecart_comp1.to_excel(writer, sheet_name='Composant Analysis', index= False)
                qte_comp1.to_excel(writer, sheet_name='Composant Analysis', startrow=0, startcol=13 , index=False)
                ecart_by_mattype1.to_excel(writer, sheet_name='MatType Analysis', index=False)
                mattypdf1.to_excel(writer, sheet_name='MatType Analysis', startrow=0, startcol=13 , index=False)
                ecart_by_Div1.to_excel(writer, sheet_name='Division Analysis', index=False)
                Divdf1.to_excel(writer, sheet_name='Division Analysis', startrow=0, startcol=13 , index=False)
                ecart_Fam1.to_excel(writer, sheet_name='Family Analysis', index= False)
                qtefam1.to_excel(writer, sheet_name='Family Analysis', startrow=0, startcol=13 , index=False)
                ecart_dec1.to_excel(writer, sheet_name='Decade Analysis', index= False)
                qtedec1.to_excel(writer, sheet_name='Decade Analysis', startrow=0, startcol=13 , index=False)
            
                workbook = writer.book
                
                worksheet = writer.sheets['MatType Analysis']
                
                chart_line_ecart = workbook.add_chart({'type': 'line'})
                chart_line_ecart.add_series({
                'categories': ['MatType Analysis', 1, 0, len(ecart_by_mattype1), 0],  
                'values':     ['MatType Analysis', 1, 1, len(ecart_by_mattype1), 1],  
                'name':       'Ecart by Mattype',
                })

                chart_line_ecart.width = 600  
                chart_line_ecart.height = 360
                worksheet.insert_chart('D1', chart_line_ecart)  
                chart_column = workbook.add_chart({'type': 'column'})
                chart_column.add_series({
                    'categories': ['MatType Analysis', 1, 0, len(mattypdf1), 0],  
                    'values':     ['MatType Analysis', 1, 14, len(mattypdf1), 14],  
                    'name':       'Cons ACT',
                })
                chart_column.add_series({
                    'values':     ['MatType Analysis', 1, 15, len(mattypdf1), 15],  
                    'name':       'Quantity Standard',
                })
                chart_column.set_title({'name': 'Comparison of Cons ACT and Quantity Standard'})
                chart_column.set_x_axis({'name': 'MatTyp'})
                chart_column.set_y_axis({'name': 'Values'})
                
                chart_column.width = 600  
                chart_column.height = 360
                worksheet.insert_chart('R1', chart_column)


                worksheet = writer.sheets['Division Analysis']
                
                chart_line_ecart2 = workbook.add_chart({'type': 'line'})
                chart_line_ecart2.add_series({
                'categories': ['Division Analysis', 1, 0, len(ecart_by_Div1), 0],  
                'values':     ['Division Analysis', 1, 1, len(ecart_by_Div1), 1],  
                'name':       'Ecart by Division',
                })

                chart_line_ecart2.width = 600  
                chart_line_ecart2.height = 360
                worksheet.insert_chart('D1', chart_line_ecart2)  
                chart_column2 = workbook.add_chart({'type': 'column'})
                chart_column2.add_series({
                    'categories': ['Division Analysis', 1, 0, len(Divdf1), 0],  
                    'values':     ['Division Analysis', 1, 14, len(Divdf1), 14],  
                    'name':       'Cons ACT',
                })
                chart_column2.add_series({
                    'values':     ['Division Analysis', 1, 15, len(Divdf1), 15],  
                    'name':       'Quantity Standard',
                })
                chart_column2.set_title({'name': 'Comparison of Cons ACT and Quantity Standard'})
                chart_column2.set_x_axis({'name': 'Division'})
                chart_column2.set_y_axis({'name': 'Values'})
                
                chart_column2.width = 600  
                chart_column2.height = 360
                worksheet.insert_chart('R1', chart_column2)

                worksheet = writer.sheets['Composant Analysis']
                
                chart_line_ecart3 = workbook.add_chart({'type': 'line'})
                chart_line_ecart3.add_series({
                'categories': ['Composant Analysis', 1, 0, len(ecart_comp1), 0],  
                'values':     ['Composant Analysis', 1, 1, len(ecart_comp1), 1],  
                'name':       'Ecart by Composant',
                })

                chart_line_ecart3.width = 600  
                chart_line_ecart3.height = 360
                worksheet.insert_chart('D1', chart_line_ecart3)  
                chart_column3 = workbook.add_chart({'type': 'column'})
                chart_column3.add_series({
                    'categories': ['Composant Analysis', 1, 0, len(qte_comp1), 0],  
                    'values':     ['Composant Analysis', 1, 14, len(qte_comp1), 14],  
                    'name':       'Cons ACT',
                })
                chart_column3.add_series({
                    'values':     ['Composant Analysis', 1, 15, len(qte_comp1), 15],  
                    'name':       'Quantity Standard',
                })
                chart_column3.set_title({'name': 'Comparison of Cons ACT and Quantity Standard'})
                chart_column3.set_x_axis({'name': 'Composant'})
                chart_column3.set_y_axis({'name': 'Values'})
                
                chart_column3.width = 600  
                chart_column3.height = 360
                worksheet.insert_chart('R1', chart_column3)

                worksheet = writer.sheets['Family Analysis']
                
                chart_line_ecart4 = workbook.add_chart({'type': 'line'})
                chart_line_ecart4.add_series({
                'categories': ['Family Analysis', 1, 0, len(ecart_Fam1), 0],  
                'values':     ['Family Analysis', 1, 1, len(ecart_Fam1), 1],  
                'name':       'Ecart by Family',
                })

                chart_line_ecart4.width = 600  
                chart_line_ecart4.height = 360
                worksheet.insert_chart('D1', chart_line_ecart4)  
                chart_column4 = workbook.add_chart({'type': 'column'})
                chart_column4.add_series({
                    'categories': ['Family Analysis', 1, 0, len(qtefam1), 0],  
                    'values':     ['Family Analysis', 1, 14, len(qtefam1), 14],  
                    'name':       'Cons ACT',
                })
                chart_column4.add_series({
                    'values':     ['Family Analysis', 1, 15, len(qtefam1), 15],  
                    'name':       'Quantity Standard',
                })
                chart_column4.set_title({'name': 'Comparison of Cons ACT and Quantity Standard'})
                chart_column4.set_x_axis({'name': 'Famille'})
                chart_column4.set_y_axis({'name': 'Values'})
                
                chart_column4.width = 600  
                chart_column4.height = 360
                worksheet.insert_chart('R1', chart_column4)

                worksheet = writer.sheets['Decade Analysis']
                
                chart_line_ecart5 = workbook.add_chart({'type': 'line'})
                chart_line_ecart5.add_series({
                'categories': ['Decade Analysis', 1, 0, len(ecart_dec1), 0],  
                'values':     ['Decade Analysis', 1, 1, len(ecart_dec1), 1],  
                'name':       'Ecart by Decade',
                })

                chart_line_ecart5.width = 600  
                chart_line_ecart5.height = 360
                worksheet.insert_chart('D1', chart_line_ecart5)  
                chart_column5 = workbook.add_chart({'type': 'column'})
                chart_column5.add_series({
                    'categories': ['Decade Analysis', 1, 0, len(qtedec1), 0],  
                    'values':     ['Decade Analysis', 1, 14, len(qtedec1), 14],  
                    'name':       'Cons ACT',
                })
                chart_column5.add_series({
                    'values':     ['Decade Analysis', 1, 15, len(qtedec1), 15],  
                    'name':       'Quantity Standard',
                })
                chart_column5.set_title({'name': 'Comparison of Cons ACT and Quantity Standard'})
                chart_column5.set_x_axis({'name': 'Decade'})
                chart_column5.set_y_axis({'name': 'Values'})
                
                chart_column5.width = 600  
                chart_column5.height = 360
                worksheet.insert_chart('R1', chart_column5)
                
              

        print("Analysis file has been created:", result_file_final)
        #########################################
        time.sleep(3)
        messagebox.showinfo("Information", "Data processing completed. Result file saved as {}".format(result_file_final))
        
    except Exception as e:
        tk.messagebox.showerror("Error", f"An error occurred: {e}")
    finally:
        processing = False
        process_button.config(state=tk.NORMAL)
        process_button.config(text="Process Data")
    progress_bar['value'] = 0
    root.update_idletasks()

    os.startfile(result_file_final)
    return merged_df
def process_data(root,input_file_entry, user_choice_entry,user_choice_combobox2, process_button,progress_bar):
    t = threading.Thread(target=process_data_thread, args=(root,input_file_entry, user_choice_entry,user_choice_combobox2, process_button,progress_bar))
    t.start() 
